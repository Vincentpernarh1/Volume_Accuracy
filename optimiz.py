import pandas as pd
from openpyxl import load_workbook
import numpy as np
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext
import threading
from tkinter import ttk
import re
import os
import time
from datetime import datetime
from PIL import Image, ImageTk
import sys

# --- Global variables ---
start_time = 0
execution_minutes = 0
output_folder_volume = None
log_widget = None
run_button = None
progress_bar = None
progress_label = None
root = None

# --- Utility Functions ---
def resource_path(relative_path):
    """Get absolute path to resource, works for dev and for PyInstaller."""
    try:
        base_path = sys._MEIPASS
    except AttributeError:
        base_path = os.path.abspath(os.path.dirname(__file__))
    return os.path.join(base_path, relative_path)

# --- Logging and Progress Update Functions (Thread-safe) ---
def log_message(message):
    """Logs a message to the console and schedules an update to the GUI."""
    print(message)
    if log_widget and root:
        root.after(0, lambda: (
            log_widget.config(state=tk.NORMAL),
            log_widget.insert(tk.END, message + "\n"),
            log_widget.config(state=tk.DISABLED),
            log_widget.see(tk.END)
        ))

def update_progress(percentage, text=""):
    """Schedules an update for the GUI progress bar and label."""
    if progress_bar and progress_label and root:
        root.after(0, lambda: (
            progress_bar.config(value=percentage),
            progress_label.config(text=f"{text} {int(percentage)}%")
        ))

# --- Core Application Logic ---
def select_and_run_process():
    global start_time, output_folder_volume, execution_minutes
    try:
        root.after(0, lambda: (
            log_widget.config(state=tk.NORMAL),
            log_widget.delete('1.0', tk.END),
            log_widget.config(state=tk.DISABLED)
        ))
        update_progress(0, "Ready")

        source_folder = filedialog.askdirectory(title="Select Folder with Input Files")
        if not source_folder:
            log_message("Process cancelled: No source folder selected.")
            return

        output_folder = filedialog.askdirectory(title="Select Output Folder")
        if not output_folder:
            log_message("Process cancelled: No output folder selected.")
            return

        output_folder_volume = os.path.join(output_folder, "Volume_Accuracy.xlsx")
        relatorio_file = os.path.join(source_folder, "Relatorio_61.xlsx")
        griglia_file = os.path.join(source_folder, "Griglia.xlsx")
        multi_de_para_file = os.path.join(source_folder, "tb_de_para.xlsx")
        
        start_time = time.time()

        missing = [f for f in [relatorio_file, griglia_file, multi_de_para_file] if not os.path.exists(f)]
        if missing:
            raise FileNotFoundError(f"Missing required files: {', '.join(os.path.basename(f) for f in missing)}")

        root.after(0, progress_bar.stop)
        main_process(relatorio_file, griglia_file, multi_de_para_file)
        
        update_progress(100, "Completed!")
        messagebox.showinfo("Success", f"Processing complete.\nExecution time: {execution_minutes:.2f} minutes")

    except Exception as e:
        log_message(f"ERROR: {e}")
        messagebox.showerror("Error", f"An error occurred:\n{str(e)}")
    finally:
        if run_button: root.after(0, lambda: run_button.config(state="normal"))
        if progress_bar: root.after(0, progress_bar.stop)

def run_process_thread():
    run_button.config(state="disabled")
    progress_bar.start()
    progress_label.config(text="Awaiting folder selection...")
    threading.Thread(target=select_and_run_process, daemon=True).start()

# --- Data Dictionaries and Cleaning Functions ---
livello_map = {"liv.0":"LL0","liv.1":"LL1","liv.2":"LL2","liv.3":"LL3","liv.4":"LL4","liv.5":"LL5","liv.6":"LL6","liv.7":"LL7","liv.8":"LL8","liv.9":"LL9","liv.10":"LL10","liv.11":"LL11","liv.12":"LL12","liv.13":"LL13"}
map_598 = {"liv.0":"Level 0","liv.1":"Level 1","liv.2":"Level 2","liv.3":"Level 3","liv.4":"Level 4","liv.5":"Level 5","liv.6":"Level 6","liv.7":"Level 7","liv.8":"Level 8","liv.9":"Level 9","liv.10":"Level 10","liv.11":"Level 11","liv.12":"Level 12","LL0":"Level 0","LL1":"Level 1","LL2":"Level 2","LL3":"Level 3","LL4":"Level 4","LL5":"Level 5","LL6":"Level 6","LL7":"Level 7","LL8":"Level 8","LL9":"Level 9","LL10":"Level 10","LL11":"Level 11","LL12":"Level 12"}
markets = ["ARGENTINA","BRASILE","MESSICO","ALTRI MERCATI","BRAZIL","MEXICO","OTHER MARKETS"]

def clean_markets():
    global markets
    markets = [i.strip().lower() for i in markets]
    return markets

def traslate_598(trans):
    match = re.search(r"liv\.?\s*\d", trans.lower())
    if match: return match.group()
    return trans

def clean_volume(val):
    if isinstance(val, str): val = val.replace(",", "").strip()
    try: return float(val)
    except (ValueError, TypeError): return None

# --- Data Processing Functions ---

def load_dataframes(File_Rela_61, File_Griglia):
    log_message("Processing: load_dataframes()")
    log_message("-> Loading Relatorio_61.xlsx and Griglia.xlsx...")
    wb_61 = load_workbook(filename=File_Rela_61, read_only=True, data_only=True)
    wb_griglia = load_workbook(filename=File_Griglia, read_only=True, data_only=True)
    if "61" not in wb_61.sheetnames:
        raise ValueError(f"Sheet '61' not found in {File_Rela_61}")
    data_61 = list(wb_61["61"].iter_rows(values_only=True))
    df_61 = pd.DataFrame(data_61[1:], columns=data_61[0]) if len(data_61) > 1 else pd.DataFrame()
    
    # --- ADDED FILTERING LOGIC ---
    log_message("-> Filtering out rows with empty 'Modelo'...")
    initial_rows = len(df_61)
    
    if 'Modelo' in df_61.columns:
        # First, drop rows where 'Modelo' is a true NaN or None
        df_61.dropna(subset=['Modelo'], inplace=True)
        
        # Next, filter out rows where 'Modelo' is a string that represents a null-like value
        # Create a boolean mask for rows to keep
        # Using .astype(str) handles cases where 'Modelo' might be numeric (e.g., 341.0)
        mask = ~df_61['Modelo'].astype(str).str.strip().str.lower().isin(['', 'none', 'nan'])
        df_61 = df_61[mask].copy()
        
        log_message(f"-> Removed {initial_rows - len(df_61)} invalid rows. Continuing with {len(df_61)} rows.")
    else:
        log_message("-> WARNING: 'Modelo' column not found. Skipping filter.")
    # --- END OF FILTERING LOGIC ---

    sheet_griglia = wb_griglia[wb_griglia.sheetnames[0]]
    data_griglia = list(sheet_griglia.iter_rows(values_only=True))
    df_griglia = pd.DataFrame(data_griglia[1:], columns=data_griglia[0]) if len(data_griglia) > 1 else pd.DataFrame()
    log_message("-> Dataframes loaded successfully.")
    return df_61, df_griglia

def clean_griglia(df_griglia):
    df_griglia["Packet_cleaned"] = df_griglia["included"].astype(str).str.upper().str.replace(r'\s+', '', regex=True)
    df_griglia["Model_cleaned"] = df_griglia["Model"].apply(
        lambda x: str(int(x)) if pd.notnull(x) and str(x).replace('.', '', 1).isdigit() and float(x).is_integer() else str(x)
    ).str.strip().str.lower()
    df_griglia["Plant_cleaned"] = df_griglia["Plant"].astype(str).str.strip()
    df_griglia["Packet"] = df_griglia["Packet"].astype(str).str.strip()
    df_griglia["Code"] = df_griglia["Code"].astype(str).str.strip()
    return df_griglia

def process_volume_table(df_61, df_griglia, field="excluded", min_col_name="Code_excluded"):
    log_message(f"Processing: process_volume_table() for '{field}' field...")
    df_griglia_clean = clean_griglia(df_griglia.copy())
    results = []
    mode = "min" if field == "included" else "max"
    for _, row_61 in df_61.iterrows():
        model_61 = str(row_61.get("Modelo")).strip()
        plant_61 = str(row_61.get("Plant")).strip()
        packet_raw = str(row_61.get(field) or "")
        filtered_griglia = df_griglia_clean[
            (df_griglia_clean["Model_cleaned"] == model_61) &
            (df_griglia_clean["Plant_cleaned"] == plant_61)
        ]
        packet_values = [val.replace(" ", "").strip().upper() for val in packet_raw.split(",") if val.strip()]
        unique_packet_values = set(packet_values)
        if filtered_griglia.empty:
            row_data = row_61.to_dict()
            row_data.update({"SINCOM": None, "volume_Head": None, "VolumeTT": None, min_col_name: 0})
            results.append(row_data)
            continue
        unique_sincom_rows = (
            filtered_griglia[['SINCOM', 'Volume Head', 'Volume TT']]
            .drop_duplicates()
            .dropna(subset=["SINCOM"])
        )
        if unique_sincom_rows.empty:
            row_data = row_61.to_dict()
            row_data.update({"SINCOM": None, "volume_Head": None, "VolumeTT": None, min_col_name: 0})
            results.append(row_data)
            continue
        for _, sincom_row in unique_sincom_rows.iterrows():
            sincom = sincom_row['SINCOM']
            volume_head = clean_volume(sincom_row.get('Volume Head'))
            volume_tt = clean_volume(sincom_row.get('Volume TT'))
            row_data = row_61.to_dict()
            row_data.update({
                "SINCOM": sincom,
                "volume_Head": volume_head,
                "VolumeTT": volume_tt,
                min_col_name: compute_volume_metric(unique_packet_values, sincom, filtered_griglia, mode=mode)
            })
            results.append(row_data)
    result_df = pd.DataFrame(results)
    preserved_cols = list(df_61.columns) + ["SINCOM", "volume_Head", "VolumeTT", min_col_name]
    for col in preserved_cols:
        if col not in result_df.columns:
            result_df[col] = np.nan if col != min_col_name else 0
    log_message(f"-> Finished processing '{field}' volume table.")
    return result_df[preserved_cols]

def compute_volume_metric(unique_packet_values, sincom, filtered_griglia, mode="max"):
    if not unique_packet_values:
        return 0
    volume_values = []
    for packet_code in unique_packet_values:
        if len(packet_code) == 1: packet_code = "00" + packet_code
        elif len(packet_code) == 2: packet_code = "0" + packet_code
        packet_code = packet_code.strip()
        matched_by_code = filtered_griglia[(filtered_griglia["SINCOM"] == sincom) & (filtered_griglia["Code"].str.contains(packet_code, na=False, regex=False))]
        matched_by_packet = filtered_griglia[(filtered_griglia["SINCOM"] == sincom) & (filtered_griglia["Code"] != packet_code) & (filtered_griglia["Packet_cleaned"].str.lower().str.startswith("pack", na=False)) & (filtered_griglia["Packet_cleaned"].str.lower().str.contains(str(packet_code).lower(), na=False, regex=False))]
        matched_packets = pd.concat([matched_by_code, matched_by_packet])
        matched_packets['Volume'] = pd.to_numeric(matched_packets['Volume'], errors='coerce')
        matched_packets = matched_packets.dropna(subset=['Volume'])
        if not matched_packets.empty:
            matched_packets = matched_packets.loc[[matched_packets['Volume'].idxmax()]]
        else:
            continue
        if matched_packets.index.empty:
            if mode == "min": return 0
            else: continue
        raw_volume = matched_packets.iloc[0].get("Volume")
        if len(matched_packets) > 1 and pd.notnull(matched_packets.iloc[1].get("Volume")):
            raw_volume = matched_packets.iloc[1].get("Volume")
        volume = clean_volume(raw_volume)
        if volume is None or volume == "":
            if mode == "min": return 0
            else: continue
        volume_values.append(volume)
    if not volume_values: return 0
    if mode == "min": return min(volume_values) if len(volume_values) == len(unique_packet_values) else 0
    else: return max(volume_values)

def main_process(file_61, file_griglia, mapping_file):
    log_message("Starting main process...")
    update_progress(5, "Loading initial data...")
    try:
        df_61, df_griglia_data = load_dataframes(file_61, file_griglia)
        # If df_61 is empty after filtering, stop the process.
        if df_61.empty:
            log_message("-> No valid data to process after filtering 'Modelo'. Stopping process.")
            return

        update_progress(10, "Processing included codes...")
        included_df = process_volume_table(df_61, df_griglia_data, field="included", min_col_name="Code_included")
        update_progress(15, "Processing excluded codes...")
        excluded_df = process_volume_table(df_61, df_griglia_data, field="excluded", min_col_name="Code_excluded")
        if excluded_df.empty and included_df.empty:
            log_message("-> Both included and excluded dataframes are empty. Exiting.")
            return
        update_progress(20, "Merging dataframes...")
        merge_keys = list(df_61.columns) + ["SINCOM", "volume_Head", "VolumeTT"]
        merge_keys = list(dict.fromkeys(merge_keys))
        merged_df = pd.merge(excluded_df, included_df, on=merge_keys, how="outer")
        extract_and_save_structured_data(merged_df, mapping_file, df_griglia_data)
        log_message("Main process completed successfully.")
    except Exception as e:
        log_message(f"-> FATAL ERROR in main_process: {e}")
        raise

def normalize_multivalues(raw):
    if pd.isna(raw) or not isinstance(raw, str): return np.nan
    raw = raw.strip()
    if raw == "": return ""
    normalized = []
    pattern = re.compile(r"(\w+)\(([^)]+)\)([+-])")
    matches = pattern.findall(raw)
    current_year, next_year = datetime.today().year % 100, (datetime.today().year % 100) + 1
    for prefix, values, sign in matches:
        if prefix == "MY":
            try:
                year_values = {int(v.strip()) for v in values.split(",") if v.strip().isdigit()}
                if year_values.intersection({current_year, next_year}): continue
            except Exception: continue
        items = [v.strip() for v in values.split(",")]
        for item in items:
            normalized.append(f"{prefix}({item}){sign}")
    return normalized if normalized else ""

def extract_and_save_structured_data(df_61, mapping_file_path, df_griglia):
    log_message("Processing: extract_and_save_structured_data()")
    update_progress(25, "Loading mapping file...")
    try:
        wb_mapping = load_workbook(mapping_file_path, data_only=True)
        df_mapping = pd.DataFrame(wb_mapping["Coded"].values)
        df_mapping.columns = df_mapping.iloc[0]
        df_mapping = df_mapping[1:].reset_index(drop=True)
        df_mapping["MultiValues"] = df_mapping["MultiValues"].astype(str).str.strip()
        mapping_dict_resp1 = df_mapping.set_index("MultiValues")["Resp.1"].to_dict()
        mapping_dict_resp2 = df_mapping.set_index("MultiValues")["Resp.2"].to_dict()
        mapping_dict_ita = df_mapping.set_index("MultiValues")["Griglia Italiano"].to_dict()
        mapping_dict_eng = df_mapping.set_index("MultiValues")["Griglia Inglês"].to_dict()
        log_message("-> Mapping dictionaries created.")
        update_progress(30, "Normalizing multivalues...")
        df_61["Multi_included"] = ""
        df_61["Multi_excluded"] = ""
        df_61["Multi_excluded_packets"] = ""
        multivalues_col = next((col for col in df_61.columns if str(col).strip().lower() == "multivalues"), None)
        if not multivalues_col:
            log_message("-> WARNING: Could not find a 'Multivalues' column.")
            return

        df_griglia = clean_griglia(df_griglia.copy())

        def process_row(row):
            raw = str(row.get(multivalues_col) or "")
            tokens = normalize_multivalues(raw)
            if not isinstance(tokens, list): tokens = []
            plant, model = str(row["Plant"]).strip(), str(row["Modelo"]).strip()
            included, excluded, packets_handled, packet_references = set(), set(), set(), set()
            for token in tokens:
                if not token: continue
                sign, token_base = token[-1], token[:-1].strip()
                translated = mapping_dict_resp1.get(token_base, token_base)
                if (plant == "FIAPE" and model in ["226", "291", "281"]) and "L" in token_base:
                    translated = livello_map.get(translated.lower().replace(" ", "").strip(), translated)
                if plant == "FIAPE" and model == "521" and token == "liv.5": translated = "LL5"
                if (plant == "FIAPE" and model in ["598", "551"]) and "L" in token_base[0]:
                    translated = map_598.get(translated.lower().replace(" ", "").strip(), translated)
                packet_ita, packet_eng = str(mapping_dict_ita.get(token_base, token_base)).strip(), str(mapping_dict_eng.get(token_base, token_base)).strip()
                if (packet_ita, packet_eng) in packets_handled: continue
                packets_handled.add((packet_ita, packet_eng))
                temp_df = df_griglia[(df_griglia["Model_cleaned"] == model) & (df_griglia["Plant_cleaned"] == plant) & (df_griglia["Packet"] == packet_ita)]
                if temp_df.index.empty:
                    temp_df = df_griglia[(df_griglia["Model_cleaned"] == model) & (df_griglia["Plant_cleaned"] == plant) & (df_griglia["Packet"] == packet_eng)]
                    packet_references.add(packet_eng)
                else: packet_references.add(packet_ita)
                griglia_multis = {v.lower().strip().replace(" ", "") for mv in temp_df["Multivalues"] if pd.notna(mv) for v in str(mv).split(",")}
                translated_tokens_from_input = set()
                for t in tokens:
                    if not t: continue
                    s, base = t[-1], t[:-1].strip()
                    trans = mapping_dict_resp1.get(base, base) if packet_ita in packet_references else mapping_dict_resp2.get(base, base)
                    if (plant == "FIAPE" and model in ["226", "291", "281"]) and "L" in base: trans = livello_map.get(trans.lower().replace(" ", "").strip(), trans)
                    if (plant == "FIASA" and model == "281") and "L" in base: trans = livello_map.get(trans.lower().replace(" ", "").strip(), trans)
                    if plant == "FIAPE" and model == "521" and base == "liv.5": trans = "LL5"
                    if (plant == "FIAPE" and model in ["598", "551"]) and "L" in base[0]:
                        trans = map_598.get(traslate_598(trans.lower().replace(" ", "").strip()), trans)
                    trans = str(trans).lower().replace(" ", "").strip()
                    translated_tokens_from_input.add(trans)
                    if s == "+": included.add(trans)
                    elif s == "-": excluded.add(trans)
                translated_tokens_lower = {token.strip().lower() for mv in translated_tokens_from_input for token in str(mv).split(",") if token.strip()}
                if sign == "+":
                    for mv in translated_tokens_from_input: included.add(mv)
                    for mv in griglia_multis:
                        if mv not in translated_tokens_lower: excluded.add(mv)
                elif sign == "-":
                    for mv in translated_tokens_from_input: excluded.add(str(mv).lower().replace(" ", "").strip())
                    for mv in griglia_multis:
                        if mv.lower() not in translated_tokens_lower: included.add(str(mv).lower().replace(" ", "").strip())
            included_strs, excluded_strs = list(map(str, included)), list(map(str, excluded))
            included_lower_set = {val.lower() for val in included_strs}
            excluded_cleaned = [val.lower() for val in excluded_strs if val.lower() not in included_lower_set]
            return (",".join(sorted(included_strs)), ",".join(sorted(excluded_cleaned)), ",".join(sorted(packet_references)))
        results = df_61.apply(process_row, axis=1, result_type='expand')
        df_61[["Multi_included", "Multi_excluded", "Multi_excluded_packets"]] = results
        log_message("-> Row processing for multivalues complete.")
        update_progress(45, "Mapping included multivalues...")
    except Exception as e:
        log_message(f"-> ERROR during data extraction: {e}")
        return None
    df_61 = map_multi_included_to_griglia(df_61, df_griglia)
    update_progress(60, "Mapping excluded multivalues...")
    df_61 = map_multi_excluded_to_griglia(df_61, df_griglia, df_mapping)
    update_progress(75, "Calculating final volumes...")
    df_61["Code_excluded"] = pd.to_numeric(df_61["Code_excluded"], errors="coerce").fillna(0)
    df_61["multi_excluded_max_volume"] = pd.to_numeric(df_61["multi_excluded_max_volume"], errors="coerce").fillna(0)
    df_61["Final_Volume_Excluded"] = np.maximum(df_61["Code_excluded"], df_61["multi_excluded_max_volume"])
    df_61["Code_included"] = pd.to_numeric(df_61["Code_included"], errors="coerce")
    df_61["multi_included_min_volume"] = pd.to_numeric(df_61["multi_included_min_volume"], errors="coerce")
    df_61["volume_Head"] = pd.to_numeric(df_61["volume_Head"], errors="coerce")
    df_61["Used_Fallback_HeadVolume"] = False
    if all(col in df_61.columns for col in ["included", "Code_included", "multi_included_min_volume"]):
        mask_include_fail = df_61["included"].notna() & (df_61["Code_included"].fillna(0) == 0)
        df_61["Final_Volume_include"] = np.where(mask_include_fail, 0, np.maximum(df_61["Code_included"].fillna(0), df_61["multi_included_min_volume"].fillna(0)))
        mask_min = (df_61["included"].notna() & (df_61["Code_included"] != 0) & (df_61["multi_included_min_volume"] != 0) & ~((df_61["multivalues"].astype(str).str.strip().str.lower().isin(["none", ""])) | (df_61["multivalues"].fillna("").astype(str).str.strip() == "")))
        df_61.loc[mask_min, "Final_Volume_include"] = np.minimum(df_61.loc[mask_min, "Code_included"], df_61.loc[mask_min, "multi_included_min_volume"])
        mask_include_condition = (df_61["Multi_included"].fillna("").astype(str).apply(lambda x: any(val.strip().upper() in markets for val in x.split(",") if val.strip())) & (df_61["included"].notna() | (df_61["Code_included"].fillna(0) == 0)))
        df_61.loc[mask_include_condition, "Final_Volume_include"] = np.maximum(df_61.loc[mask_include_condition, "Code_included"].fillna(0), df_61.loc[mask_include_condition, "multi_included_min_volume"].fillna(0))
        mask_multi_include_condition = ((df_61["included"].astype(str).str.strip().str.lower().isin(["none", ""])) | (df_61["included"].fillna("").astype(str).str.strip() == "")) & (df_61["multi_included_min_volume"] != 0)
        df_61.loc[mask_multi_include_condition, "Final_Volume_include"] = np.maximum(df_61.loc[mask_multi_include_condition, "Code_included"].fillna(0), df_61.loc[mask_multi_include_condition, "multi_included_min_volume"].fillna(0))
        if all(col in df_61.columns for col in ["excluded", "Code_excluded", "volume_Head"]):
            fallback_mask = ((df_61["Final_Volume_include"] == 0) & (df_61["Multi_included"].fillna("").str.strip().str.lower().isin(["none", ""])) & (df_61["included"].fillna("").str.strip().str.lower().isin(["none", "", "0"])) & df_61["excluded"].notna() & df_61["volume_Head"].notna())
            df_61.loc[fallback_mask, "Final_Volume_include"] = df_61.loc[fallback_mask, "volume_Head"] - df_61.loc[fallback_mask, "Code_excluded"]
            df_61.loc[fallback_mask, "Used_Fallback_HeadVolume"] = True
    df_61["Final_Volume_include"] = pd.to_numeric(df_61["Final_Volume_include"], errors="coerce")
    df_61["Final_Volume_Excluded"] = pd.to_numeric(df_61.get("Final_Volume_Excluded", 0), errors="coerce").fillna(0)
    df_61["Volume_Mix"] = np.where(df_61["Used_Fallback_HeadVolume"], df_61["Final_Volume_include"], np.maximum(df_61["Final_Volume_include"] - df_61["Final_Volume_Excluded"], 0))
    update_progress(90, "Aggregating results...")
    df_61["Volume_Mix"] = pd.to_numeric(df_61["Volume_Mix"], errors="coerce")
    final_columns = [col for col in ["Modelo", "PN", "Plant", "multivalues", "included", "Multi_excluded", "excluded", "Unique_Key", "VolumeTT"] if col in df_61.columns]
    df_61 = df_61.groupby(final_columns, dropna=False)[["Volume_Mix"]].sum().reset_index()
    all_empty_mask = ((df_61["multivalues"].fillna("").str.strip().str.lower().isin(["", "none"])) | ((df_61["multivalues"].fillna("").str.strip() != "") & (df_61["Multi_excluded"].fillna("").str.strip() == ""))) & (df_61["included"].fillna("").str.strip() == "") & (df_61["excluded"].fillna("").str.strip() == "")
    df_61.loc[all_empty_mask, "Volume_Mix"] = df_61.loc[all_empty_mask, "VolumeTT"]
    df_61.drop(columns=["Multi_excluded"], inplace=True, errors='ignore')
    df_61['VolumeTT'] = pd.to_numeric(df_61['VolumeTT'], errors='coerce')
    df_61["Mix"] = df_61["Volume_Mix"].divide(df_61["VolumeTT"]).fillna(0)
    update_progress(95, "Saving final file...")
    df_61.to_excel(output_folder_volume, index=False)
    global execution_minutes
    end_time = time.time()
    execution_minutes = (end_time - start_time) / 60
    log_message(f"-> Final file saved to {output_folder_volume}")
    log_message(f"-> Total execution time: {execution_minutes:.2f} minutes.")
    
def map_multi_included_to_griglia(df_flattened, df_griglia):
    log_message("Processing: map_multi_included_to_griglia()")
    df_flattened_copy = df_flattened.copy()
    df_flattened_copy["Modelo"] = df_flattened_copy["Modelo"].astype(str).str.strip().str.lower()
    df_griglia["Model_cleaned"] = df_griglia["Model"].astype(str).str.strip().str.lower()
    df_griglia["Plant_cleaned"] = df_griglia["Plant"].astype(str).str.strip()
    df_griglia["Multivalues_cleaned_list"] = df_griglia["Multivalues"].astype(str).str.lower().str.replace(r'\s+', '', regex=True).str.split(",")
    min_volumes = []
    for _, row in df_flattened_copy.iterrows():
        model, plant, sincom = row["Modelo"], str(row["Plant"]).strip(), str(row["SINCOM"]).strip()
        multi_included = str(row.get("Multi_included", "")).lower()
        included_tokens = [val.strip().replace(" ", "") for val in multi_included.split(",") if val.strip()]
        matched_volumes = []
        filtered_griglia = df_griglia[(df_griglia["Model_cleaned"] == model) & (df_griglia["Plant_cleaned"] == plant)]
        for token in included_tokens:
            if plant == "FIAPE" and model in ["226", "291", "281"]: token = livello_map.get(token, token)
            if plant == "FIAPE" and model == "521" and token == "liv.5": token = "LL5"
            if token in [m.lower() for m in markets]:
                token_griglia_rows = filtered_griglia[filtered_griglia["included"].str.lower().str.replace(" ", "").str.contains(token, na=False, regex=False)]
            elif len(token) <= 3:
                token_griglia_rows = filtered_griglia[filtered_griglia["Multivalues_cleaned_list"].apply(lambda x: token in x if isinstance(x, list) else False)]
            else:
                token_griglia_rows = filtered_griglia[filtered_griglia["Multivalues"].str.lower().str.replace(" ", "").str.contains(token, na=False, regex=False)]
            for _, gr_row in token_griglia_rows.iterrows():
                if str(gr_row["SINCOM"]).strip() == sincom:
                    try: matched_volumes.append(float(gr_row.get("Volume Head", "")))
                    except (ValueError, TypeError): continue
        min_volumes.append(min(matched_volumes) if matched_volumes else np.nan)
    df_flattened_copy["multi_included_min_volume"] = min_volumes
    log_message("-> Mapped multi-included volumes.")
    return df_flattened_copy

def map_multi_excluded_to_griglia(df_flattened, griglia_path, df_mapping):
    log_message("Processing: map_multi_excluded_to_griglia()")
    griglia = griglia_path.copy()
    griglia["Model_cleaned"] = griglia["Model"].astype(str).str.strip().str.lower()
    griglia["Plant_cleaned"] = griglia["Plant"].astype(str).str.strip()
    griglia["Multivalue_token"] = griglia["Multivalues"].astype(str).str.lower().str.split(",").apply(lambda tokens: [t.strip() for t in tokens])
    griglia_exploded = griglia.explode("Multivalue_token").reset_index(drop=True)
    df_flattened_copy = df_flattened.copy()
    df_flattened_copy["multi_excluded_max_volume"] = 0
    def apply_row(row):
        model, plant, sincom = str(row["Modelo"]).strip().lower(), str(row["Plant"]).strip(), str(row["SINCOM"]).strip()
        multi_excluded_tokens = [t.strip().replace(" ", "").lower() for t in str(row["Multi_excluded"]).split(",") if t.strip()]
        multi_excluded_packets = [p.strip() for p in str(row["Multi_excluded_packets"]).split(",") if p.strip()] if pd.notna(row["Multi_excluded_packets"]) else []
        if not multi_excluded_tokens: return 0
        filtered = griglia_exploded[(griglia_exploded["Model_cleaned"] == model) & (griglia_exploded["Plant_cleaned"] == plant) & (griglia_exploded["SINCOM"] == sincom)]
        if multi_excluded_packets: filtered = filtered[filtered["Packet"].isin(multi_excluded_packets)]
        multivalue_tokens_in_filtered = filtered["Multivalue_token"].astype(str).str.replace(" ", "").str.lower().tolist()
        if any(token in multivalue_tokens_in_filtered for token in multi_excluded_tokens):
            numeric_volumes = pd.to_numeric(filtered["Volume Head"], errors='coerce').dropna().tolist()
            return max(numeric_volumes) if numeric_volumes else 0
        return 0
    mask = df_flattened_copy["Multi_excluded"].fillna("").str.lower().isin(["nan", "none", ""]) == False
    if mask.any():
        df_flattened_copy.loc[mask, "multi_excluded_max_volume"] = df_flattened_copy.loc[mask].apply(apply_row, axis=1)
    log_message("-> Mapped multi-excluded volumes.")
    return df_flattened_copy

def create_gui():
    global root, progress_bar, progress_label, log_widget, run_button
    root = tk.Tk()
    root.title("Volume Accuracy Processor")
    root.geometry("600x500")
    root.configure(bg="#f0f0f0")
    try:
        stellantis_logo_path = resource_path("assets/Vlc_img.png")
        img_stellantis_logo = Image.open(stellantis_logo_path)
        img_stellantis_logo = img_stellantis_logo.resize((530, 40), Image.Resampling.LANCZOS)
        photo_img_stellantis_logo = ImageTk.PhotoImage(img_stellantis_logo)
        root.image = photo_img_stellantis_logo
        tk.Label(root, image=photo_img_stellantis_logo, bg="#f0f0f0").pack(pady=10)
    except Exception as e:
        print(f"Warning: Could not load image. {e}")
        tk.Label(root, text="STELLANTIS", font=("Helvetica", 16, "bold"), bg="#f0f0f0").pack(pady=10)
    tk.Label(root, text="MIX GENERATION", font=("Helvetica", 14, "bold"), bg="#f0f0f0").pack()
    progress_bar = ttk.Progressbar(root, orient="horizontal", length=400, mode="determinate")
    progress_bar.pack(pady=10)
    progress_label = tk.Label(root, text="", font=("Helvetica", 10), bg="#f0f0f0")
    progress_label.pack()
    log_frame = tk.Frame(root, bg="#f0f0f0")
    log_frame.pack(pady=10, fill="both", expand=True, padx=20)
    log_widget = scrolledtext.ScrolledText(log_frame, state=tk.DISABLED, height=10, wrap=tk.WORD, font=("Courier New", 9))
    log_widget.pack(fill="both", expand=True)
    run_button = tk.Button(root, text="Select Folders and Run", command=run_process_thread, height=2, width=30, bg="#4CAF50", fg="white", font=("Helvetica", 12, "bold"))
    run_button.pack(pady=20)
    root.mainloop()

if __name__ == "__main__":
    create_gui()