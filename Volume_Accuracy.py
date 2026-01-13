import pandas as pd
from openpyxl import load_workbook
import numpy as np
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext
import threading
from tkinter import ttk
import re
import os
import time
from datetime import datetime
from PIL import Image, ImageTk
import sys
import polars as pl
import Volume_Parten

# --- Global variables ---
start_time = 0
execution_minutes = 0
output_folder_volume = None
log_widget = None
run_button = None
progress_bar = None
progress_label = None
root = None

# --- Utility Functions ---
def resource_path(relative_path):
    """Get absolute path to resource, works for dev and for PyInstaller."""
    try:
        base_path = sys._MEIPASS
    except AttributeError:
        base_path = os.path.abspath(os.path.dirname(__file__))
    return os.path.join(base_path, relative_path)

# --- Logging and Progress Update Functions (Thread-safe) ---
def log_message(message):
    """Logs a message to the console and schedules an update to the GUI."""
    print(message)
    if log_widget and root:
        root.after(0, lambda: (
            log_widget.config(state=tk.NORMAL),
            log_widget.insert(tk.END, message + "\n"),
            log_widget.config(state=tk.DISABLED),
            log_widget.see(tk.END)
        ))

def update_progress(percentage, text=""):
    """Schedules an update for the GUI progress bar and label."""
    if progress_bar and progress_label and root:
        root.after(0, lambda: (
            progress_bar.config(value=percentage),
            progress_label.config(text=f"{text} {int(percentage)}%")
        ))

def choose_calculation_type(parent):
    """Custom dialog to choose calculation type."""
    dialog = tk.Toplevel(parent)
    dialog.title("Choose Calculation Type")
    dialog.geometry("350x200")
    dialog.resizable(False, False)
    dialog.transient(parent)
    dialog.grab_set()
    
    # Center the dialog
    screen_width = parent.winfo_screenwidth()
    screen_height = parent.winfo_screenheight()
    x = (screen_width - 350) // 2
    y = (screen_height - 200) // 2
    dialog.geometry(f"350x200+{x}+{y}")
    
    # STELLANTIS styling
    stellantis_blue = "#003DA5"
    stellantis_orange = "#FF6600"
    
    tk.Label(dialog, text="Select Calculation Type", font=("Segoe UI", 14, "bold"), fg=stellantis_blue).pack(pady=10)
    
    var = tk.StringVar(value="normal")
    
    frame = tk.Frame(dialog)
    frame.pack(pady=10)
    
    tk.Radiobutton(frame, text="Normal (Relatorio_61)", variable=var, value="normal", font=("Segoe UI", 11)).pack(anchor=tk.W, pady=5)
    tk.Radiobutton(frame, text="Parten", variable=var, value="parten", font=("Segoe UI", 11)).pack(anchor=tk.W, pady=5)
    
    def on_ok():
        dialog.result = var.get()
        dialog.destroy()
    
    button_frame = tk.Frame(dialog)
    button_frame.pack(pady=10)
    tk.Button(button_frame, text="Select", command=on_ok, bg=stellantis_blue, fg="white", font=("Segoe UI", 10, "bold"), width=10).pack()
    
    parent.wait_window(dialog)
    return getattr(dialog, 'result', 'normal')

# --- Core Application Logic ---
def select_and_run_process():
    global start_time, output_folder_volume, execution_minutes
    try:
        root.after(0, lambda: (
            log_widget.config(state=tk.NORMAL),
            log_widget.delete('1.0', tk.END),
            log_widget.config(state=tk.DISABLED)
        ))
        update_progress(0, "Ready")

        source_folder = filedialog.askdirectory(title="Select Folder with Input Files")
        if not source_folder:
            log_message("Process cancelled: No source folder selected.")
            return

        output_folder = filedialog.askdirectory(title="Select Output Folder")
        if not output_folder:
            log_message("Process cancelled: No output folder selected.")
            return

        # Prompt for calculation type
        calculation_type = choose_calculation_type(root)

        if calculation_type == 'normal':
            output_folder_volume = os.path.join(output_folder, "Volume_Accuracy.xlsx")
            relatorio_file = os.path.join(source_folder, "Relatorio_61.xlsx")
            griglia_file = os.path.join(source_folder, "Griglia.xlsx")
            multi_de_para_file = os.path.join(source_folder, "tb_de_para.xlsx")
            
            start_time = time.time()

            missing = [f for f in [relatorio_file, griglia_file, multi_de_para_file] if not os.path.exists(f)]
            if missing:
                raise FileNotFoundError(f"Missing required files: {', '.join(os.path.basename(f) for f in missing)}")

            root.after(0, progress_bar.stop)
            main_process(relatorio_file, griglia_file, multi_de_para_file)
            
            update_progress(100, "Completed!")
            messagebox.showinfo("Success", f"Normal processing complete.\nExecution time: {execution_minutes:.2f} minutes")
        else:
            # Call Parten process
            Volume_Parten.run_parten_process(source_folder, output_folder, log_widget, progress_bar, progress_label, run_button, root)

    except Exception as e:
        log_message(f"ERROR: {e}")
        messagebox.showerror("Error", f"An error occurred:\n{str(e)}")
    finally:
        if run_button: root.after(0, lambda: run_button.config(state="normal"))
        if progress_bar: root.after(0, progress_bar.stop)

def run_process_thread():
    run_button.config(state="disabled")
    progress_bar.start()
    progress_label.config(text="Awaiting folder selection...")
    threading.Thread(target=select_and_run_process, daemon=True).start()

# --- Data Dictionaries and Cleaning Functions ---
livello_map = {"liv.0":"LL0","liv.1":"LL1","liv.2":"LL2","liv.3":"LL3","liv.4":"LL4","liv.5":"LL5","liv.6":"LL6","liv.7":"LL7","liv.8":"LL8","liv.9":"LL9","liv.10":"LL10","liv.11":"LL11","liv.12":"LL12","liv.13":"LL13"}
map_598 = {"liv.0":"Level 0","liv.1":"Level 1","liv.2":"Level 2","liv.3":"Level 3","liv.4":"Level 4","liv.5":"Level 5","liv.6":"Level 6","liv.7":"Level 7","liv.8":"Level 8","liv.9":"Level 9","liv.10":"Level 10","liv.11":"Level 11","liv.12":"Level 12","LL0":"Level 0","LL1":"Level 1","LL2":"Level 2","LL3":"Level 3","LL4":"Level 4","LL5":"Level 5","LL6":"Level 6","LL7":"Level 7","LL8":"Level 8","LL9":"Level 9","LL10":"Level 10","LL11":"Level 11","LL12":"Level 12"}
markets = ["ARGENTINA","BRASILE","MESSICO","ALTRI MERCATI","BRAZIL","MEXICO","OTHER MARKETS"]

def clean_markets():
    global markets
    markets = [i.strip().lower() for i in markets]
    return markets

def traslate_598(trans):
    match = re.search(r"liv\.?\s*\d", trans.lower())
    if match: return match.group()
    return trans

def clean_volume(val):
    if isinstance(val, str): val = val.replace(",", "").strip()
    try: return float(val)
    except (ValueError, TypeError): return None

# --- Data Processing Functions ---

def load_dataframes(File_Rela_61, File_Griglia):
    log_message("Processing: load_dataframes()")
    log_message("-> Loading Relatorio_61.xlsx and Griglia.xlsx...")

    wb_61 = load_workbook(filename=File_Rela_61, read_only=True, data_only=True)
    wb_griglia = load_workbook(filename=File_Griglia, read_only=True, data_only=True)

    if "61" not in wb_61.sheetnames:
        raise ValueError(f"Sheet '61' not found in {File_Rela_61}")

    data_61 = list(wb_61["61"].iter_rows(values_only=True))
    df_61 = pd.DataFrame(data_61[1:], columns=data_61[0]).astype(str) if len(data_61) > 1 else pd.DataFrame()

    # --- NEW: Clean up the '.0' from specific columns ---
    if not df_61.empty:
        cols_to_clean = ['Modelo', 'PN'] # Add any other numeric columns here
        for col in cols_to_clean:
            if col in df_61.columns:
                # This safely removes '.0' only if it's at the very end of the string
                df_61[col] = df_61[col].str.removesuffix('.0')
    # --- End of new section ---

    log_message("-> Filtering out rows with empty 'Modelo'...")
    if not df_61.empty and 'Modelo' in df_61.columns:
        initial_rows = len(df_61)
        mask = ~df_61['Modelo'].str.strip().str.lower().isin(['', 'none', 'nan'])
        df_61 = df_61[mask].copy()
       
    elif 'Modelo' not in df_61.columns:
        log_message("-> WARNING: 'Modelo' column not found. Skipping filter.")

    sheet_griglia = wb_griglia[wb_griglia.sheetnames[0]]
    data_griglia = list(sheet_griglia.iter_rows(values_only=True))
    df_griglia = pd.DataFrame(data_griglia[1:], columns=data_griglia[0]).astype(str)

    # Convert to Polars DataFrames
    df_61_pl = pl.from_pandas(df_61)
    df_griglia_pl = pl.from_pandas(df_griglia)

    log_message("✅ DataFrames loaded and prepared successfully.")
    return df_61_pl, df_griglia_pl

def clean_griglia(df_griglia):
    df_griglia = df_griglia.with_columns([
        pl.col("included").cast(pl.Utf8).str.to_uppercase().str.replace_all(r'\s+', '', literal=False).alias("Packet_cleaned"),
        pl.col("Model").cast(pl.Utf8).map_elements(
            lambda x: str(int(float(x))) if x.replace('.', '', 1).isdigit() and float(x).is_integer() else str(x),
            return_dtype=pl.Utf8
        ).str.strip_chars().str.to_lowercase().alias("Model_cleaned"),
        pl.col("Plant").cast(pl.Utf8).str.strip_chars().alias("Plant_cleaned"),
        pl.col("Packet").cast(pl.Utf8).str.strip_chars().alias("Packet"),
        pl.col("Code").cast(pl.Utf8).str.strip_chars().alias("Code")
    ])
    return df_griglia

def process_volume_table(df_61, df_griglia, field="excluded", min_col_name="Code_excluded"):
    log_message(f"Processing: process_volume_table() for '{field}' field...")
    
    # Using the trusted logic from your original working code
    df_griglia_clean = clean_griglia(df_griglia.clone())
    mode = "min" if field == "included" else "max"

    # Process each row in df_61
    results = []
    for row in df_61.iter_rows(named=True):
        model_61 = str(row.get("Modelo")).strip()
        plant_61 = str(row.get("Plant")).strip()
        packet_raw = str(row.get(field) or "")

        filtered_griglia = df_griglia_clean.filter(
            (pl.col("Model_cleaned") == model_61) &
            (pl.col("Plant_cleaned") == plant_61)
        )

        packet_values = [val.replace(" ", "").strip().upper() for val in packet_raw.split(",") if val.strip()]
        unique_packet_values = set(packet_values)

        if filtered_griglia.is_empty():
            row_data = row.copy()
            row_data.update({
                "SINCOM": None, "volume_Head": None, "VolumeTT": None, min_col_name: 0
            })
            results.append(row_data)
            continue

        # Get unique SINCOM rows
        unique_sincom_rows = (
            filtered_griglia
            .select(["SINCOM", "Volume Head", "Volume TT"])
            .unique()
            .filter(pl.col("SINCOM").is_not_null())
        )

        # If there are no SINCOMs for this Model/Plant, create a single null entry and move on
        if unique_sincom_rows.is_empty():
            row_data = row.copy()
            row_data.update({
                "SINCOM": None, "volume_Head": None, "VolumeTT": None, min_col_name: 0
            })
            results.append(row_data)
            continue

        # If there are SINCOMs, process each one
        for sincom_row in unique_sincom_rows.iter_rows(named=True):
            sincom = sincom_row['SINCOM']
            volume_head = clean_volume(sincom_row.get('Volume Head'))
            volume_tt = clean_volume(sincom_row.get('Volume TT'))
            row_data = row.copy()
            row_data.update({
                "SINCOM": sincom,
                "volume_Head": volume_head,
                "VolumeTT": volume_tt,
                min_col_name: compute_volume_metric(unique_packet_values, sincom, filtered_griglia, mode=mode)
            })
            results.append(row_data)
    
    result_df = pl.DataFrame(results)
    preserved_cols = list(df_61.columns) + ["SINCOM", "volume_Head", "VolumeTT", min_col_name]
    # Ensure all columns are present, even if empty
    for col in preserved_cols:
        if col not in result_df.columns:
            if col == min_col_name:
                result_df = result_df.with_columns(pl.lit(0).alias(col))
            else:
                result_df = result_df.with_columns(pl.lit(None).alias(col))
            
    log_message(f"-> Finished processing '{field}' volume table.")
    return result_df.select(preserved_cols)

def compute_volume_metric(unique_packet_values, sincom, filtered_griglia, mode="max"):
    if not unique_packet_values:
        return 0

    volume_values = []
    for packet_code in unique_packet_values:
        if len(packet_code) == 1:
            packet_code = "00" + packet_code
        elif len(packet_code) == 2:
            packet_code = "0" + packet_code
        packet_code = packet_code.strip()


        # --- Helper function to normalize the search term ---
        def normalize_search_term(code_str):
           
            try:
                # Convert to float first to handle decimals like '45.0', then to int
                return str(int(float(code_str)))
            except (ValueError, TypeError):
                # If conversion fails, it's not a number, so return the original string
                return code_str

        
        search_term = normalize_search_term(packet_code)

        # 2. Use the cleaned search_term to filter the 'Code' column.
        #    We ensure the 'Code' column is treated as a string to use .str.contains()
        matched_by_code = filtered_griglia.filter(
            (pl.col("SINCOM") == sincom) &
            (pl.col("Code").cast(pl.Utf8).str.contains(search_term, literal=True))
        )
                
        # Match by packet name (e.g., in 'included' column of griglia)
        matched_by_packet = filtered_griglia.filter(
            (pl.col("SINCOM") == sincom) &
            (pl.col("Code") != packet_code) &
            (pl.col("Packet_cleaned").str.to_lowercase().str.starts_with("pack")) &
            (pl.col("Packet_cleaned").str.to_lowercase().str.contains(str(packet_code).lower(), literal=True))
        )
        
        matched_packets = pl.concat([matched_by_code, matched_by_packet])

        matched_packets = matched_packets.with_columns(
            pl.col("Volume").cast(pl.Utf8).str.replace(",", "").str.strip_chars().cast(pl.Float64, strict=False).alias("Volume")
        ).filter(pl.col("Volume").is_not_null())

        if not matched_packets.is_empty():
            # Important: find the single best match for this packet_code
            matched_packets = matched_packets.sort("Volume", descending=True).head(1)
           
        else:
            continue
        
        if matched_packets.is_empty():
            if mode == "min":
                return 0  # In "min" mode, all packets MUST have a match
            else:
                continue # In "max" mode, we can skip missing packets

        raw_volume = matched_packets.select("Volume").to_series().to_list()[0]
        
        volume = clean_volume(raw_volume)

        if volume is None or volume == "":
            if mode == "min":
                return 0
            else:
                continue
        volume_values.append(volume)
        
    if not volume_values:
        return 0
    if mode == "min":
        return min(volume_values) if len(volume_values) == (len(unique_packet_values)) else 0
    else:
        # For "max" mode, return the max of whatever we found
        return max(volume_values)
    

def main_process(file_61, file_griglia, mapping_file):
    log_message("Starting main process...")
    update_progress(5, "Loading initial data...")
    try:
        df_61, df_griglia_data = load_dataframes(file_61, file_griglia)
        if df_61.is_empty():
            log_message("-> No valid data to process after filtering 'Modelo'. Stopping process.")
            return
        update_progress(10, "Processing included codes...")
        included_df = process_volume_table(df_61, df_griglia_data, field="included", min_col_name="Code_included")
        update_progress(15, "Processing excluded codes...")
        excluded_df = process_volume_table(df_61, df_griglia_data, field="excluded", min_col_name="Code_excluded")
        if excluded_df.is_empty() and included_df.is_empty():
            log_message("-> Both included and excluded dataframes are empty. Exiting.")
            return
        update_progress(20, "Merging dataframes...")
        merge_keys = list(df_61.columns) + ["SINCOM", "volume_Head", "VolumeTT"]
        merge_keys = list(dict.fromkeys(merge_keys))
        merged_df = excluded_df.join(included_df, on=merge_keys, how="outer")
        extract_and_save_structured_data(merged_df, mapping_file, df_griglia_data)
        log_message("Main process completed successfully.")
    except Exception as e:
        log_message(f"-> FATAL ERROR in main_process: {e}")
        raise

def normalize_multivalues(raw):
    if pd.isna(raw) or not isinstance(raw, str): return np.nan
    raw = raw.strip()
    if raw == "": return ""
    normalized = []
    pattern = re.compile(r"(\w+)\(([^)]+)\)([+-])")
    matches = pattern.findall(raw)
    current_year, next_year = datetime.today().year % 100, (datetime.today().year % 100) + 1
    for prefix, values, sign in matches:
        if prefix == "MY":
            try:
                year_values = {int(v.strip()) for v in values.split(",") if v.strip().isdigit()}
                if year_values.intersection({current_year, next_year}): continue
            except Exception: continue
        items = [v.strip() for v in values.split(",")]
        for item in items:
            normalized.append(f"{prefix}({item}){sign}")
    return normalized if normalized else ""

def extract_and_save_structured_data(df_61, mapping_file_path, df_griglia):
    log_message("Processing: extract_and_save_structured_data()")
    update_progress(25, "Loading mapping file...")
    try:
        wb_mapping = load_workbook(mapping_file_path, data_only=True)
        df_mapping = pd.DataFrame(wb_mapping["Coded"].values)
        df_mapping.columns = df_mapping.iloc[0]
        df_mapping = df_mapping[1:].reset_index(drop=True)
        df_mapping["MultiValues"] = df_mapping["MultiValues"].astype(str).str.strip()
        mapping_dict_resp1 = df_mapping.set_index("MultiValues")["Resp.1"].to_dict()
        mapping_dict_resp2 = df_mapping.set_index("MultiValues")["Resp.2"].to_dict()
        mapping_dict_ita = df_mapping.set_index("MultiValues")["Griglia Italiano"].to_dict()
        mapping_dict_eng = df_mapping.set_index("MultiValues")["Griglia Inglês"].to_dict()
        log_message("-> Mapping dictionaries created.")
        update_progress(30, "Normalizing multivalues...")
        df_61 = df_61.with_columns([
            pl.lit("").alias("Multi_included"),
            pl.lit("").alias("Multi_excluded"),
            pl.lit("").alias("Multi_excluded_packets")
        ])
        multivalues_col = next((col for col in df_61.columns if str(col).strip().lower() == "multivalues"), None)
        if not multivalues_col:
            log_message("-> WARNING: Could not find a 'Multivalues' column.")
            return

        df_griglia = clean_griglia(df_griglia.clone())

        # Convert to pandas for row-wise processing (this part is complex and iterative)
        df_61_pd = df_61.to_pandas()
        df_griglia_pd = df_griglia.to_pandas()

        def process_row(row):
            raw = str(row.get(multivalues_col) or "").strip()
            # --- OPTIMIZATION ---
            # If multivalues is empty or 'none', skip all processing
            if not raw or raw.lower() == 'none':
                return ("", "", "")
            tokens = normalize_multivalues(raw)
            if not isinstance(tokens, list): tokens = []
            plant, model = str(row["Plant"]).strip(), str(row["Modelo"]).strip()
            included, excluded, packets_handled, packet_references = set(), set(), set(), set()
            for token in tokens:
                if not token: continue
                sign, token_base = token[-1], token[:-1].strip()
                translated = mapping_dict_resp1.get(token_base, token_base)
                if (plant == "FIAPE" and model in ["226", "291", "281"]) and "L" in token_base:
                    translated = livello_map.get(translated.lower().replace(" ", "").strip(), translated)
                if plant == "FIAPE" and model == "521" and token == "liv.5": translated = "LL5"
                if (plant == "FIAPE" and model in ["598", "551"]) and "L" in token_base[0]:
                    translated = map_598.get(translated.lower().replace(" ", "").strip(), translated)
                packet_ita, packet_eng = str(mapping_dict_ita.get(token_base, token_base)).strip(), str(mapping_dict_eng.get(token_base, token_base)).strip()
                if (packet_ita, packet_eng) in packets_handled: continue
                packets_handled.add((packet_ita, packet_eng))
                temp_df = df_griglia_pd[(df_griglia_pd["Model_cleaned"] == model) & (df_griglia_pd["Plant_cleaned"] == plant) & (df_griglia_pd["Packet"] == packet_ita)]
                
                
                if temp_df.index.empty:
                    temp_df = df_griglia_pd[(df_griglia_pd["Model_cleaned"] == model) & (df_griglia_pd["Plant_cleaned"] == plant) & (df_griglia_pd["Packet"] == packet_eng)]
                    packet_references.add(packet_eng)
                else: packet_references.add(packet_ita)
                griglia_multis = {v.lower().strip().replace(" ", "") for mv in temp_df["Multivalues"] if pd.notna(mv) for v in str(mv).split(",")}
                translated_tokens_from_input = set()
                
                # print(griglia_multis)
                
                for t in tokens:
                    if not t: continue
                    s, base = t[-1], t[:-1].strip()
                    trans = mapping_dict_resp1.get(base, base) if packet_ita in packet_references else mapping_dict_resp2.get(base, base)
                    if (plant == "FIAPE" and model in ["226", "291", "281"]) and "L" in base: trans = livello_map.get(trans.lower().replace(" ", "").strip(), trans)
                   
                    
                    if (plant == "FIASA" and model == "281") and "L" in base: trans = livello_map.get(trans.lower().replace(" ", "").strip(), trans)
                    if plant == "FIAPE" and model == "521" and base == "liv.5": trans = "LL5"
                    if (plant == "FIAPE" and model in ["598", "551"]) and "L" in base[0]:
                        trans = map_598.get(traslate_598(trans.lower().replace(" ", "").strip()), trans)
                    trans = str(trans).lower().replace(" ", "").strip()
                   
                    translated_tokens_from_input.add(trans)
                    if s == "+": included.add(trans)
                    elif s == "-": excluded.add(trans)

                translated_tokens_lower = {token.strip().lower() for mv in translated_tokens_from_input for token in str(mv).split(",") if token.strip()}
                if sign == "+":
                    for mv in griglia_multis:
                        if mv not in translated_tokens_lower: excluded.add(mv)
                elif sign == "-":
                    for mv in griglia_multis:
                        if mv.lower() not in translated_tokens_lower: included.add(str(mv).lower().replace(" ", "").strip())
            
            included_strs, excluded_strs = list(map(str, included)), list(map(str, excluded))
            included_lower_set = {val.lower() for val in included_strs}
            excluded_cleaned = [val.lower() for val in excluded_strs if val.lower() not in included_lower_set]
            return (",".join(sorted(included_strs)), ",".join(sorted(excluded_cleaned)), ",".join(sorted(packet_references)))
        results = df_61_pd.apply(process_row, axis=1, result_type='expand')
        df_61_pd[["Multi_included", "Multi_excluded", "Multi_excluded_packets"]] = results
        # Convert back to Polars
        df_61 = pl.from_pandas(df_61_pd)
        log_message("-> Row processing for multivalues complete.")
        update_progress(45, "Mapping included multivalues...")
    except Exception as e:
        log_message(f"-> ERROR during data extraction: {e}")
        return None
    
    # Convert to pandas for Excel output and complex calculations
    df_61_pd = df_61.to_pandas()
    df_61_pd = map_multi_included_to_griglia(df_61_pd, df_griglia.to_pandas())
    
    update_progress(60, "Mapping excluded multivalues...")
    df_61_pd = map_multi_excluded_to_griglia(df_61_pd, df_griglia.to_pandas(), df_mapping)
    update_progress(75, "Calculating final volumes...")
    df_61_pd["Code_excluded"] = pd.to_numeric(df_61_pd["Code_excluded"], errors="coerce").fillna(0)
    df_61_pd["multi_excluded_max_volume"] = pd.to_numeric(df_61_pd["multi_excluded_max_volume"], errors="coerce").fillna(0)
    df_61_pd["Final_Volume_Excluded"] = np.maximum(df_61_pd["Code_excluded"], df_61_pd["multi_excluded_max_volume"])
    df_61_pd["Code_included"] = pd.to_numeric(df_61_pd["Code_included"], errors="coerce")
    df_61_pd["multi_included_min_volume"] = pd.to_numeric(df_61_pd["multi_included_min_volume"], errors="coerce").fillna(0)
    df_61_pd["volume_Head"] = pd.to_numeric(df_61_pd["volume_Head"], errors="coerce")
    df_61_pd["Used_Fallback_HeadVolume"] = False

    # df_61_pd.to_excel("After_Included.xlsx", index=False)

    if all(col in df_61_pd.columns for col in ["included", "Code_included", "multi_included_min_volume"]):

               
        mask_include_fail = df_61_pd["included"].notna() & (df_61_pd["Code_included"].fillna(0) == 0)
        df_61_pd["Final_Volume_include"] = np.where(mask_include_fail, 0, np.maximum(df_61_pd["Code_included"].fillna(0), df_61_pd["multi_included_min_volume"].fillna(0)))
        

        mask_min = (df_61_pd["included"].notna() & (df_61_pd["Code_included"] != 0) & (df_61_pd["multi_included_min_volume"] != 0) & ~((df_61_pd["multivalues"].astype(str).str.strip().str.lower().isin(["none", ""])) | (df_61_pd["multivalues"].fillna("").astype(str).str.strip() == "")) | (df_61_pd["multivalues"].astype(str).str.strip().str.lower().isin(["MY(26)+","MY(26)-","MY(27)+","MY(27)-","MY(28)+","MY(28)-","MY(29)+","MY(29)-","MY(30)+"])))
        
        df_61_pd.loc[mask_min, "Final_Volume_include"] = np.minimum(df_61_pd.loc[mask_min, "Code_included"], df_61_pd.loc[mask_min, "multi_included_min_volume"])
        
        

        #This is remove for the sake of markets being recalculated and also
        
        # mask_include_condition = (df_61["Multi_included"].fillna("").astype(str).apply(lambda x: any(val.strip().upper() in markets for val in x.split(",") if val.strip())) & (df_61["included"].notna() | (df_61["Code_included"].fillna(0) == 0)))
        # df_61.loc[mask_include_condition, "Final_Volume_include"] = np.maximum(df_61.loc[mask_include_condition, "Code_included"].fillna(0), df_61.loc[mask_include_condition, "multi_included_min_volume"].fillna(0))
       
        mask_multi_include_condition = ((df_61_pd["included"].astype(str).str.strip().str.lower().isin(["none", ""])) | (df_61_pd["included"].fillna("").astype(str).str.strip() == "")) & (df_61_pd["multi_included_min_volume"] != 0)
        df_61_pd.loc[mask_multi_include_condition, "Final_Volume_include"] = np.maximum(df_61_pd.loc[mask_multi_include_condition, "Code_included"].fillna(0), df_61_pd.loc[mask_multi_include_condition, "multi_included_min_volume"].fillna(0))
        if all(col in df_61_pd.columns for col in ["excluded", "Code_excluded", "volume_Head"]):
            fallback_mask = ((df_61_pd["Final_Volume_include"] == 0) & (df_61_pd["Multi_included"].fillna("").str.strip().str.lower().isin(["none", ""])) & (df_61_pd["included"].fillna("").str.strip().str.lower().isin(["none", "", "0"])) & df_61_pd["excluded"].notna() & df_61_pd["volume_Head"].notna())
            df_61_pd.loc[fallback_mask, "Final_Volume_include"] = df_61_pd.loc[fallback_mask, "volume_Head"] - df_61_pd.loc[fallback_mask, "Code_excluded"]
            df_61_pd.loc[fallback_mask, "Used_Fallback_HeadVolume"] = True
   
    df_61_pd["Final_Volume_include"] = pd.to_numeric(df_61_pd["Final_Volume_include"], errors="coerce")
    df_61_pd["Final_Volume_Excluded"] = pd.to_numeric(df_61_pd.get("Final_Volume_Excluded", 0), errors="coerce").fillna(0)
    df_61_pd["Volume_Mix"] = np.where(df_61_pd["Used_Fallback_HeadVolume"], df_61_pd["Final_Volume_include"], np.maximum(df_61_pd["Final_Volume_include"] - df_61_pd["Final_Volume_Excluded"], 0))
    update_progress(90, "Aggregating results...")
    df_61_pd["Volume_Mix"] = pd.to_numeric(df_61_pd["Volume_Mix"], errors="coerce")
    final_columns = [col for col in ["Modelo", "PN", "Plant", "multivalues", "included", "Multi_excluded", "excluded", "Unique_Key", "VolumeTT"] if col in df_61_pd.columns]
    df_61_pd = df_61_pd.groupby(final_columns, dropna=False)[["Volume_Mix"]].sum().reset_index()
    all_empty_mask = ((df_61_pd["multivalues"].fillna("").str.strip().str.lower().isin(["", "none", "nan","MY(26)+","MY(26)-","MY(27)+","MY(27)-","MY(28)+"])) | ((df_61_pd["multivalues"].fillna("").str.strip() != "") & (df_61_pd["Multi_excluded"].fillna("").str.strip() == ""))) & (df_61_pd["included"].fillna("").str.strip() == "") & (df_61_pd["excluded"].fillna("").str.strip() == "")
    df_61_pd.loc[all_empty_mask, "Volume_Mix"] = df_61_pd.loc[all_empty_mask, "VolumeTT"]
    df_61_pd.drop(columns=["Multi_excluded"], inplace=True, errors='ignore')
    df_61_pd['VolumeTT'] = pd.to_numeric(df_61_pd['VolumeTT'], errors='coerce')
    df_61_pd["Mix"] = df_61_pd["Volume_Mix"].divide(df_61_pd["VolumeTT"]).fillna(0)
    update_progress(95, "Saving final file...")
    df_61_pd.to_excel(output_folder_volume, index=False)
    global execution_minutes
    end_time = time.time()
    execution_minutes = (end_time - start_time) / 60
    log_message(f"-> Final file saved to {output_folder_volume}")
    log_message(f"-> Total execution time: {execution_minutes:.2f} minutes.")
    

def map_multi_included_to_griglia(df_flattened, df_griglia):
    log_message("Processing: map_multi_included_to_griglia()")
    df_flattened_copy = df_flattened.copy()
    df_flattened_copy["Modelo"] = df_flattened_copy["Modelo"].astype(str).str.strip().str.lower()
    df_griglia["Model_cleaned"] = df_griglia["Model"].astype(str).str.strip().str.lower()
    df_griglia["Plant_cleaned"] = df_griglia["Plant"].astype(str).str.strip()
    df_griglia["Multivalues_cleaned_list"] = df_griglia["Multivalues"].astype(str).str.lower().str.replace(r'\s+', '', regex=True).str.split(",")
    min_volumes = []
    for _, row in df_flattened_copy.iterrows():
        model, plant, sincom = row["Modelo"], str(row["Plant"]).strip(), str(row["SINCOM"]).strip()
        multi_included = str(row.get("Multi_included", "")).lower()
        included_tokens = [val.strip().replace(" ", "") for val in multi_included.split(",") if val.strip()]
        matched_volumes = []
        filtered_griglia = df_griglia[(df_griglia["Model_cleaned"] == model) & (df_griglia["Plant_cleaned"] == plant)]

        for token in included_tokens:
            if plant == "FIAPE" and model in ["226", "291", "281"]: token = livello_map.get(token, token)
            if plant == "FIAPE" and model == "521" and token == "liv.5": token = "LL5"
            if token in [m.lower() for m in markets]:
                token_griglia_rows = filtered_griglia[filtered_griglia["included"].str.lower().str.replace(" ", "").str.contains(token, na=False, regex=False)]
            elif len(token) <= 3:
                token_griglia_rows = filtered_griglia[filtered_griglia["Multivalues_cleaned_list"].apply(lambda x: token in x if isinstance(x, list) else False)]
            else:
                token_griglia_rows = filtered_griglia[filtered_griglia["Multivalues"].str.lower().str.replace(" ", "").str.contains(token, na=False, regex=False)]
           
            for _, gr_row in token_griglia_rows.iterrows():
                if str(gr_row["SINCOM"]).strip() == sincom:
                    if token in [m.lower() for m in markets]: 
                        try: matched_volumes.append(float(gr_row.get("Volume", "")))
                        except (ValueError, TypeError): continue
                    else :
                        try: matched_volumes.append(float(gr_row.get("Volume Head", "")))
                        except (ValueError, TypeError): continue

        min_volumes.append(min(matched_volumes) if matched_volumes else np.nan)
    df_flattened_copy["multi_included_min_volume"] = min_volumes
    log_message("-> Mapped multi-included volumes.")
    return df_flattened_copy

def map_multi_excluded_to_griglia(df_flattened, griglia_path, df_mapping):
    log_message("Processing: map_multi_excluded_to_griglia()")
    griglia = griglia_path.copy()
    griglia["Model_cleaned"] = griglia["Model"].astype(str).str.strip().str.lower()
    griglia["Plant_cleaned"] = griglia["Plant"].astype(str).str.strip()
    griglia["Multivalue_token"] = griglia["Multivalues"].astype(str).str.lower().str.split(",").apply(lambda tokens: [t.strip() for t in tokens])
    griglia_exploded = griglia.explode("Multivalue_token").reset_index(drop=True)
    df_flattened_copy = df_flattened.copy()
    df_flattened_copy["multi_excluded_max_volume"] = 0
    def apply_row(row):
        model, plant, sincom = str(row["Modelo"]).strip().lower(), str(row["Plant"]).strip(), str(row["SINCOM"]).strip()
        multi_excluded_tokens = [t.strip().replace(" ", "").lower() for t in str(row["Multi_excluded"]).split(",") if t.strip()]
        multi_excluded_packets = [p.strip() for p in str(row["Multi_excluded_packets"]).split(",") if p.strip()] if pd.notna(row["Multi_excluded_packets"]) else []
        if not multi_excluded_tokens: return 0
        filtered = griglia_exploded[(griglia_exploded["Model_cleaned"] == model) & (griglia_exploded["Plant_cleaned"] == plant) & (griglia_exploded["SINCOM"] == sincom)]
        if multi_excluded_packets: filtered = filtered[filtered["Packet"].isin(multi_excluded_packets)]
        multivalue_tokens_in_filtered = filtered["Multivalue_token"].astype(str).str.replace(" ", "").str.lower().tolist()
        
        if any(token in multivalue_tokens_in_filtered for token in multi_excluded_tokens):
            numeric_volumes = pd.to_numeric(filtered["Volume Head"], errors='coerce').dropna().tolist()
            return max(numeric_volumes) if numeric_volumes else 0
        return 0
    mask = df_flattened_copy["Multi_excluded"].fillna("").str.lower().isin(["nan", "none", ""]) == False
    if mask.any():
        df_flattened_copy.loc[mask, "multi_excluded_max_volume"] = df_flattened_copy.loc[mask].apply(apply_row, axis=1)
    log_message("-> Mapped multi-excluded volumes.")
    return df_flattened_copy

def create_gui():
    global root, progress_bar, progress_label, log_widget, run_button
    root = tk.Tk()
    root.title("Volume Accuracy Processor")
    # Center the window on the screen
    screen_width = root.winfo_screenwidth()
    screen_height = root.winfo_screenheight()
    width = 700
    height = 550
    x = (screen_width - width) // 2
    y = (screen_height - height) // 4
    root.geometry(f"{width}x{height}+{x}+{y}")
    root.resizable(True, True)
    
    # STELLANTIS Colors
    stellantis_blue = "#003DA5"
    stellantis_orange = "#FF6600"
    dhl_yellow = "#FFCC00"
    
    # Set modern color scheme with STELLANTIS theme
    style = ttk.Style()
    style.theme_use('clam')
    
    # Configure button style with STELLANTIS blue
    style.configure('TButton', background=stellantis_blue, foreground="white", relief="flat", padding=6, font=("Segoe UI", 10, "bold"))
    style.map('TButton', background=[('active', stellantis_orange)])
    
    # Configure progressbar with STELLANTIS colors
    style.configure('TProgressbar', background=stellantis_blue, troughcolor='#E8E8E8', bordercolor='#CCCCCC', lightcolor=stellantis_orange, darkcolor=stellantis_blue)
    
    # Configure labels with theme colors
    style.configure('Title.TLabel', font=("Segoe UI", 16, "bold"), foreground=stellantis_blue)

    # --- Main container ---
    container = tk.Frame(root, bg="white")
    container.pack(fill=tk.BOTH, expand=True)

    # --- Header with STELLANTIS accent ---
    header_frame = tk.Frame(container, bg=stellantis_blue, height=60)
    header_frame.pack(fill=tk.X, padx=0, pady=0)
    header_frame.pack_propagate(False)
    
    # Load and place the STELLANTIS logo image to fill the header
    img_stellantis_logo_original = None
    try:
        stellantis_logo_path = resource_path("assets/Vlc_img.png")
        img_stellantis_logo_original = Image.open(stellantis_logo_path)
        img_stellantis_logo = img_stellantis_logo_original.resize((700, 60), Image.Resampling.LANCZOS)
        photo_img_stellantis_logo = ImageTk.PhotoImage(img_stellantis_logo)
        root.image = photo_img_stellantis_logo
        logo_label = tk.Label(header_frame, image=photo_img_stellantis_logo, bg=stellantis_blue)
        logo_label.pack(fill=tk.BOTH, expand=True)
    except Exception as e:
        print(f"Warning: Could not load image. {e}")
        logo_label = tk.Label(header_frame, text="STELLANTIS", font=("Segoe UI", 16, "bold"), fg="white", bg=stellantis_blue)
        logo_label.pack(fill=tk.BOTH, expand=True)
    
    # Function to resize image on window resize
    def resize_image(event=None):
        if img_stellantis_logo_original and hasattr(root, 'image'):
            width = root.winfo_width()
            if width > 0:
                resized = img_stellantis_logo_original.resize((width, 60), Image.Resampling.LANCZOS)
                photo = ImageTk.PhotoImage(resized)
                root.image = photo
                logo_label.config(image=photo)
    
    root.bind('<Configure>', resize_image)

    # --- Main content frame ---
    main_frame = ttk.Frame(container, padding="13")
    main_frame.pack(fill=tk.BOTH, expand=True)

    # Title
    title_main = ttk.Label(main_frame, text="MIX GENERATION", style='Title.TLabel')
    title_main.pack(pady=(0, 10))

    # Status section
    progress_label = ttk.Label(main_frame, text="Ready to start. Click 'Select Folders and Run'.", font=("Segoe UI", 11), foreground=stellantis_blue)
    progress_label.pack(pady=(2, 5), padx=1, fill=tk.X)

    # Progress bar with accent color
    progress_bar = ttk.Progressbar(main_frame, orient='horizontal', length=400, mode='determinate')
    progress_bar.pack(pady=10, padx=5, fill=tk.X)

    # Button section with modern styling
    button_frame = ttk.Frame(main_frame)
    button_frame.pack(pady=4, fill=tk.X)
    
    run_button = ttk.Button(button_frame, text="▶ Select Folders and Run", command=run_process_thread, style='TButton')
    run_button.pack(anchor=tk.CENTER, padx=5)
    
    # Log section with accent
    log_frame = ttk.LabelFrame(main_frame, text="📋 Activity Log", padding="13")
    log_frame.pack(pady=0, padx=2, fill=tk.BOTH, expand=True)
    
    log_widget = scrolledtext.ScrolledText(log_frame, wrap=tk.WORD, width=80, height=10, font=("Consolas", 11), bg="#F5F5F5", fg="#333333", state=tk.DISABLED)
    log_widget.pack(fill=tk.BOTH, expand=True)
    
    # Footer section with STELLANTIS branding
    footer_frame = tk.Frame(container, bg=stellantis_blue, height=34)
    footer_frame.pack(fill=tk.X, padx=0, pady=0, side=tk.BOTTOM)
    footer_frame.pack_propagate(False)
    
    # Left side - STELLANTIS branding
    left_footer = tk.Frame(footer_frame, bg=stellantis_blue)
    left_footer.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=15, pady=10)
    
    stellantis_label = tk.Label(left_footer, text="🏢 STELLANTIS", font=("Segoe UI", 11, "bold"), fg=stellantis_orange, bg=stellantis_blue)
    stellantis_label.pack(side=tk.LEFT, padx=1)
    
    # Right side - Developer credit
    right_footer = tk.Frame(footer_frame, bg=stellantis_blue)
    right_footer.pack(side=tk.RIGHT, padx=15, pady=10)
    
    footer_label = tk.Label(right_footer, text="Developed by: Vincent Pernarh", font=("Segoe UI", 9), fg="white", bg=stellantis_blue)
    footer_label.pack(anchor="e")

    root.mainloop()

if __name__ == "__main__":
    create_gui()